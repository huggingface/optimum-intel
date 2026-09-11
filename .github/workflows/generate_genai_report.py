"""
Generate a markdown report summarizing the latest workflow run for test_openvino_genai_aipc.yml.

Fetches job data and logs from the GitHub Actions API, parses pytest results per model,
and produces a markdown table with models in rows, MTL/LNL columns with CPU/GPU/NPU subcolumns.

Requirements:
    pip install requests
    pip install openpyxl  # only needed for --xlsx

Usage:
    # Set GITHUB_TOKEN environment variable (needs repo/actions read access)
    # Or have `gh` CLI authenticated (token will be auto-detected)
    set GITHUB_TOKEN=ghp_...
    python generate_genai_report.py

    # Optionally specify a run ID:
    python generate_genai_report.py --run-id 29916107736

    # Save to file (markdown):
    python generate_genai_report.py --output optimum-genai-test-report.md

    # Generate HTML report:
    python generate_genai_report.py --html --output optimum-genai-test-report.html

    # Generate Excel report (single sheet, transformers version as column):
    python generate_genai_report.py --xlsx --output optimum-genai-test-report.xlsx
"""

import argparse
import os
import re
import subprocess
import sys
from collections import defaultdict

import requests


OWNER = os.environ.get("GITHUB_REPOSITORY_OWNER", "huggingface")
REPO_NAME = os.environ.get("GITHUB_REPOSITORY", "huggingface/optimum-intel").split("/")[-1]
WORKFLOW_FILE = "test_openvino_genai_aipc.yml"
API_BASE = f"https://api.github.com/repos/{OWNER}/{REPO_NAME}"


def get_headers():
    token = os.environ.get("GITHUB_TOKEN")
    if not token:
        # Try to get token from gh CLI
        try:
            result = subprocess.run(["gh", "auth", "token"], capture_output=True, text=True, check=True)
            token = result.stdout.strip()
        except (subprocess.CalledProcessError, FileNotFoundError):
            pass
    if not token:
        print(
            "ERROR: GITHUB_TOKEN environment variable is required (or authenticate via `gh auth login`).",
            file=sys.stderr,
        )
        sys.exit(1)
    return {
        "Authorization": f"token {token}",
        "Accept": "application/vnd.github.v3+json",
    }


def get_latest_run(run_id=None):
    """Get the latest completed workflow run, or a specific run by ID."""
    headers = get_headers()
    if run_id:
        url = f"{API_BASE}/actions/runs/{run_id}"
        resp = requests.get(url, headers=headers)
        resp.raise_for_status()
        return resp.json()
    else:
        url = f"{API_BASE}/actions/workflows/{WORKFLOW_FILE}/runs"
        params = {"per_page": 10, "status": "completed"}
        resp = requests.get(url, headers=headers, params=params)
        resp.raise_for_status()
        runs = resp.json()["workflow_runs"]
        if not runs:
            print("No completed workflow runs found.", file=sys.stderr)
            sys.exit(1)
        return runs[0]


def get_jobs(run_id):
    """Get all jobs for a workflow run."""
    headers = get_headers()
    jobs = []
    page = 1
    while True:
        url = f"{API_BASE}/actions/runs/{run_id}/jobs"
        resp = requests.get(url, headers=headers, params={"per_page": 100, "page": page})
        resp.raise_for_status()
        data = resp.json()
        jobs.extend(data["jobs"])
        if len(jobs) >= data["total_count"]:
            break
        page += 1
    return jobs


def download_job_log(job_id):
    """Download log for a single job. Returns the log text or None on failure."""
    headers = get_headers()
    url = f"{API_BASE}/actions/jobs/{job_id}/logs"
    resp = requests.get(url, headers=headers, allow_redirects=True)
    if resp.status_code == 200:
        return resp.text
    return None


def parse_job_name(job_name):
    """Parse job name like 'test (4.57.6, GPU, MTL)' into (version, device, runner)."""
    match = re.match(r"test\s*\(([^,]+),\s*([^,]+),\s*([^)]+)\)", job_name)
    if match:
        return match.group(1).strip(), match.group(2).strip(), match.group(3).strip()
    return None, None, None


def parse_versions_from_log(log_text):
    """Extract package versions from the 'Verify imports' step output.

    Looks for lines like:
        openvino_genai 2026.3.0.0-3272-e9c0cf80d50
        openvino 2026.3.0-22446-ce03a15d415-releases/2026/3
        transformers 4.57.6
        torch 2.13.0+cpu

    Returns dict of {package_name: version_string}.
    """
    versions = {}
    timestamp_re = re.compile(r"^\d{4}-\d{2}-\d{2}T[\d:.]+Z\s?")
    for line in log_text.split("\n"):
        stripped = timestamp_re.sub("", line).strip()
        for pkg in ("openvino_genai", "openvino", "transformers", "torch"):
            if stripped.startswith(f"{pkg} "):
                versions[pkg] = stripped[len(pkg) + 1 :].strip()
                break
    return versions


def parse_test_results_from_log(log_text):
    """Parse pytest output from a GitHub Actions job log to extract per-test results.

    The log format is:
        2026-07-22T11:35:00.6366112Z tests/openvino/test_genai.py::ClassName::method_name
        ... (optional log lines) ...
        2026-07-22T11:35:06.8379394Z PASSED                                   [ xx%]

    Or inline (for some tests):
        2026-07-22T... tests/openvino/test_genai.py::Class::method SKIPPED [xx%]

    Under pytest-xdist the result is reported by the controller instead:
        2026-07-22T... [gw0] [ xx%] PASSED tests/openvino/test_genai.py::Class::method

    A test whose worker died from a native crash (access violation) is reported as:
        2026-07-22T... worker 'gw0' crashed while running 'tests/openvino/test_genai.py::Class::method'

    Tests that were collected but never reported a status (the run was aborted before reaching
    them) are returned as 'not_run' so they are not confused with passing tests.

    Returns dict of {(class_name, test_method): status} where status is 'passed', 'failed',
    'skipped', 'error', or 'not_run'.
    """
    results = {}
    collected = set()
    crashed = set()
    lines = log_text.split("\n")

    # Strip timestamp prefix from lines: "2026-07-22T11:35:00.6366112Z "
    timestamp_re = re.compile(r"^\d{4}-\d{2}-\d{2}T[\d:.]+Z\s?")

    current_test = None  # (class_name, method_name)

    for line in lines:
        stripped = timestamp_re.sub("", line).strip()

        # A crashed worker takes the whole process down, so the status is recorded separately
        # and applied last: xdist also reports the test as a plain failure.
        crash_match = re.search(
            r"worker '\w+' crashed while running '(?:tests/openvino/)?test_genai\.py::(\w+)::(\w+)'",
            stripped,
        )
        if crash_match:
            crashed.add((crash_match.group(1), crash_match.group(2)))
            continue

        # Check for a status-first result line, emitted both by pytest-xdist progress
        # ("[gw0] [ xx%] PASSED test_genai.py::Class::method") and by the short test
        # summary that -r produces ("FAILED test_genai.py::Class::method").
        status_first_match = re.match(
            r"(?:\[gw\d+\]\s*)?(?:\[\s*\d+%\]\s*)?(PASSED|FAILED|SKIPPED|ERROR)\s+"
            r"(?:tests/openvino/)?test_genai\.py::(\w+)::(\w+)",
            stripped,
        )
        if status_first_match:
            status = status_first_match.group(1).lower()
            results[(status_first_match.group(2), status_first_match.group(3))] = status
            current_test = None
            continue

        # Check for inline result: "test_genai.py::Class::method PASSED/FAILED/SKIPPED [xx%]"
        inline_match = re.match(
            r"(?:tests/openvino/)?test_genai\.py::(\w+)::(\w+)\s+(PASSED|FAILED|SKIPPED)",
            stripped,
        )
        if inline_match:
            class_name = inline_match.group(1)
            method = inline_match.group(2)
            status = inline_match.group(3).lower()
            results[(class_name, method)] = status
            current_test = None
            continue

        # Check for test name line: "tests/openvino/test_genai.py::ClassName::method_name"
        test_match = re.match(
            r"(?:tests/openvino/)?test_genai\.py::(\w+)::(\w+)\s*$",
            stripped,
        )
        if test_match:
            current_test = (test_match.group(1), test_match.group(2))
            collected.add(current_test)
            continue

        # Check for standalone result line: "PASSED  [ xx%]" or "FAILED  [ xx%]" or "SKIPPED (reason)  [ xx%]"
        # These always end with a percentage indicator like "[ xx%]"
        result_match = re.match(r"(PASSED|FAILED|SKIPPED)\b", stripped)
        if result_match and current_test and re.search(r"\[\s*\d+%\]", stripped):
            status = result_match.group(1).lower()
            results[current_test] = status
            current_test = None
            continue

        # Check for crash/access violation (indicates current test errored)
        if current_test and "fatal exception" in stripped.lower():
            results[current_test] = "error"
            current_test = None
            continue

    for test in crashed:
        results[test] = "error"

    for test in collected:
        results.setdefault(test, "not_run")

    return results


def extract_model_from_test_name(test_method):
    """Extract model architecture name from test method like 'test_compare_outputs_00_gpt2'.

    parameterized adds a numeric prefix: test_compare_outputs_00_gpt_bigcode, test_compare_outputs_01_bloom, etc.
    Returns None if the test method doesn't match the expected pattern.
    """
    match = re.match(r"test_compare_outputs_(?:vlm_)?\d+_(.+)", test_method)
    if match:
        return match.group(1)
    return None


def extract_test_category(class_name):
    """Map test class name to a category."""
    mapping = {
        "LLMPipelineTestCase": "LLM",
        "VLMPipelineTestCase": "VLM",
        "Speech2TextPipelineTestCase": "Speech2Text",
        "Text2SpeechPipelineTestCase": "Text2Speech",
        "LLMPipelineWithEagle3TestCase": "Eagle3",
        "VLMPipelineWithEagle3TestCase": "Eagle3-VLM",
    }
    return mapping.get(class_name, class_name)


def build_report(run_info, jobs, job_logs):
    """Build the markdown report from jobs and their logs.

    Args:
        run_info: Workflow run metadata from API
        jobs: List of job objects from API
        job_logs: Dict of {job_id: log_text}
    """
    # Structure: results[version][(category, model_name)] = {(runner, device): status}
    results = defaultdict(lambda: defaultdict(dict))
    job_conclusions = {}
    # Track versions per transformers version (from any job with that version)
    version_info = {}  # {transformers_version: {pkg: version_str}}

    for job in jobs:
        job_name = job["name"]
        version, device, runner = parse_job_name(job_name)
        if not version:
            continue

        conclusion = job.get("conclusion", "unknown")
        job_conclusions[(version, device, runner)] = conclusion

        log_content = job_logs.get(job["id"])
        if log_content:
            # Extract package versions (once per transformers version is enough)
            if version not in version_info:
                versions = parse_versions_from_log(log_content)
                if versions:
                    version_info[version] = versions

            test_results = parse_test_results_from_log(log_content)
            for (class_name, test_method), status in test_results.items():
                category = extract_test_category(class_name)
                model_name = extract_model_from_test_name(test_method)
                if model_name is None:
                    continue  # Skip unparseable test names
                model_key = (category, model_name)
                results[version][model_key][(runner, device)] = status

    # If no per-test results found, fall back to job-level conclusions
    if not any(results.values()):
        print(
            "WARNING: Could not parse per-test results from logs. Using job-level conclusions only.", file=sys.stderr
        )
        for job in jobs:
            job_name = job["name"]
            version, device, runner = parse_job_name(job_name)
            if not version:
                continue
            conclusion = job.get("conclusion", "unknown")
            results[version][("Job", "all_tests")][(runner, device)] = conclusion

    # Generate markdown
    lines = []
    run_url = run_info.get("html_url", "")
    run_number = run_info.get("run_number", "")
    run_status = run_info.get("conclusion", run_info.get("status", "unknown"))
    head_branch = run_info.get("head_branch", "")
    created_at = run_info.get("created_at", "")

    lines.append("# OpenVINO GenAI AIPC Test Report")
    lines.append("")
    lines.append(f"**Workflow Run:** [#{run_number}]({run_url})")
    lines.append(f"**Branch:** `{head_branch}`")
    lines.append(f"**Status:** {run_status}")
    lines.append(f"**Date:** {created_at}")
    lines.append("")

    # Package versions table
    if version_info:
        lines.append("### Package Versions")
        lines.append("")
        lines.append("| Transformers | OpenVINO | OpenVINO GenAI | PyTorch |")
        lines.append("|:---:|:---:|:---:|:---:|")
        for tv in sorted(version_info.keys()):
            vi = version_info[tv]
            ov = vi.get("openvino", "-")
            ov_genai = vi.get("openvino_genai", "-")
            torch_ver = vi.get("torch", "-")
            tf_ver = vi.get("transformers", tv)
            lines.append(f"| {tf_ver} | {ov} | {ov_genai} | {torch_ver} |")
        lines.append("")

    # Status emoji mapping
    status_symbols = {
        "passed": "\u2705",
        "success": "\u2705",
        "failed": "\u274c",
        "failure": "\u274c",
        "skipped": "\u23ed\ufe0f",
        "error": "\u26a0\ufe0f",
        "not_run": "\u2b1c",
        "cancelled": "\u23f9\ufe0f",
        "unknown": "\u2753",
    }

    runners = ["MTL", "LNL"]
    devices = ["CPU", "GPU", "NPU"]

    for version in sorted(results.keys()):
        lines.append(f"## Transformers {version}")
        lines.append("")

        # Table header with sub-columns
        header = "| Category | Model |"
        separator = "|----------|-------|"
        for runner in runners:
            for device in devices:
                header += f" {runner}/{device} |"
                separator += ":---:|"

        lines.append(header)
        lines.append(separator)

        # Group by category, sorted
        model_keys = sorted(results[version].keys(), key=lambda x: (x[0], x[1]))

        for category, model_name in model_keys:
            row = f"| {category} | {model_name} |"
            for runner in runners:
                for device in devices:
                    status = results[version][(category, model_name)].get((runner, device), "-")
                    symbol = status_symbols.get(status, status)
                    row += f" {symbol} |"
            lines.append(row)

        lines.append("")

    # Legend
    lines.append("## Legend")
    lines.append("")
    lines.append("| Symbol | Meaning |")
    lines.append("|--------|---------|")
    lines.append("| \u2705 | Passed |")
    lines.append("| \u274c | Failed |")
    lines.append("| \u23ed\ufe0f | Skipped |")
    lines.append("| \u26a0\ufe0f | Error (crash/access violation) |")
    lines.append("| \u2b1c | Not run (collected, but the run ended first) |")
    lines.append("| \u23f9\ufe0f | Cancelled |")
    lines.append("| - | Not run / No data |")
    lines.append("")

    # Job-level summary
    lines.append("## Job-Level Summary")
    lines.append("")
    header = "| Version |"
    separator = "|---------|"
    for runner in runners:
        for device in devices:
            header += f" {runner}/{device} |"
            separator += ":---:|"
    lines.append(header)
    lines.append(separator)

    for version in sorted({v for v, _, _ in job_conclusions.keys()}):
        row = f"| {version} |"
        for runner in runners:
            for device in devices:
                conclusion = job_conclusions.get((version, device, runner), "-")
                symbol = status_symbols.get(conclusion, conclusion)
                row += f" {symbol} |"
        lines.append(row)
    lines.append("")

    # Direct links to jobs
    lines.append("<details><summary>Job Links</summary>")
    lines.append("")
    lines.append("| Job | Conclusion |")
    lines.append("|-----|-----------|")
    for job in sorted(jobs, key=lambda j: j["name"]):
        job_name = job["name"]
        conclusion = job.get("conclusion", "unknown")
        symbol = status_symbols.get(conclusion, conclusion)
        job_url = job.get("html_url", "")
        lines.append(f"| [{job_name}]({job_url}) | {symbol} {conclusion} |")
    lines.append("")
    lines.append("</details>")
    lines.append("")

    return "\n".join(lines)


def build_xlsx_report(run_info, jobs, job_logs, output_path):
    """Build an Excel report with a single sheet combining all transformers versions.

    Columns: Transformers Version | Category | Model | MTL/CPU | MTL/GPU | MTL/NPU | LNL/CPU | LNL/GPU | LNL/NPU
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        print("ERROR: openpyxl is required for --xlsx. Install with: pip install openpyxl", file=sys.stderr)
        sys.exit(1)

    # Parse results (same logic as build_report)
    results = defaultdict(lambda: defaultdict(dict))
    # Map (version, runner, device) -> job URL for hyperlinks
    job_urls = {}

    for job in jobs:
        job_name = job["name"]
        version, device, runner = parse_job_name(job_name)
        if not version:
            continue

        job_url = job.get("html_url", "")
        if job_url:
            job_urls[(version, runner, device)] = job_url

        log_content = job_logs.get(job["id"])
        if log_content:
            test_results = parse_test_results_from_log(log_content)
            for (class_name, test_method), status in test_results.items():
                category = extract_test_category(class_name)
                model_name = extract_model_from_test_name(test_method)
                if model_name is None:
                    continue
                model_key = (category, model_name)
                results[version][model_key][(runner, device)] = status

    runners = ["MTL", "LNL"]
    devices = ["CPU", "GPU", "NPU"]

    wb = Workbook()
    ws = wb.active
    ws.title = "GenAI Test Results"

    # Header row
    headers = ["Transformers Version", "Category", "Model"]
    for runner in runners:
        for device in devices:
            headers.append(f"{runner}/{device}")
    ws.append(headers)

    # Style header
    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="F6F8FA", end_color="F6F8FA", fill_type="solid")
    for col_idx, cell in enumerate(ws[1], 1):
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Status fill colors
    status_fills = {
        "passed": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
        "failed": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
        "skipped": PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid"),
        "error": PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid"),
        "not_run": PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid"),
    }

    # Build column order for status columns: [(runner, device), ...]
    col_order = [(runner, device) for runner in runners for device in devices]

    # Data rows - sorted by version, then category, then model
    row_versions = []  # track version per data row for hyperlink lookup
    for version in sorted(results.keys()):
        model_keys = sorted(results[version].keys(), key=lambda x: (x[0], x[1]))
        for category, model_name in model_keys:
            row = [version, category, model_name]
            for runner, device in col_order:
                status = results[version][(category, model_name)].get((runner, device), "")
                row.append(status)
            ws.append(row)
            row_versions.append(version)

    # Apply status fills, center alignment, and hyperlinks to data cells
    link_font = Font(underline="single", color="0563C1")
    for row_idx in range(2, ws.max_row + 1):
        version = row_versions[row_idx - 2]
        # Left-align text columns
        for col_idx in range(1, 4):
            ws.cell(row=row_idx, column=col_idx).alignment = Alignment(horizontal="left")
        # Center, color, and hyperlink status columns
        for col_offset, (runner, device) in enumerate(col_order):
            col_idx = 4 + col_offset
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.alignment = Alignment(horizontal="center")
            fill = status_fills.get(cell.value)
            if fill:
                cell.fill = fill
            # Add hyperlink to the job run
            url = job_urls.get((version, runner, device))
            if url and cell.value:
                cell.hyperlink = url
                cell.font = link_font

    # Auto-fit column widths
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_length + 2, 30)

    # Add metadata in a freeze pane note: freeze the header row
    ws.freeze_panes = "A2"

    wb.save(output_path)


def markdown_to_html(md_text):
    """Convert markdown report to a self-contained HTML page.

    Handles: headings, bold, inline code, tables, links, <details> blocks, and paragraphs.
    Uses GitHub-like styling with no external dependencies.
    """
    html_lines = []

    html_lines.append(
        """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>OpenVINO GenAI AIPC Test Report</title>
<style>
body { font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Helvetica, Arial, sans-serif; max-width: 1400px; margin: 0 auto; padding: 20px; color: #24292f; background: #fff; }
h1 { border-bottom: 1px solid #d0d7de; padding-bottom: 8px; }
h2 { border-bottom: 1px solid #d0d7de; padding-bottom: 6px; margin-top: 24px; }
h3 { margin-top: 20px; }
table { border-collapse: collapse; margin: 12px 0; width: auto; }
th, td { border: 1px solid #d0d7de; padding: 6px 12px; text-align: center; }
th { background: #f6f8fa; font-weight: 600; }
td:first-child, td:nth-child(2) { text-align: left; }
a { color: #0969da; text-decoration: none; }
a:hover { text-decoration: underline; }
code { background: #f6f8fa; padding: 2px 6px; border-radius: 3px; font-size: 90%; }
details { margin: 12px 0; }
summary { cursor: pointer; font-weight: 600; }
p { margin: 4px 0; }
</style>
</head>
<body>
"""
    )

    in_table = False
    lines = md_text.split("\n")

    for line in lines:
        # Details blocks (pass through)
        if line.strip().startswith("<details"):
            html_lines.append(line)
            continue
        if line.strip() == "</details>":
            html_lines.append(line)
            continue
        if line.strip().startswith("<summary"):
            html_lines.append(line)
            continue

        # Table
        if line.strip().startswith("|"):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            # Skip separator rows
            if all(re.match(r"^[-:]+$", c) for c in cells):
                continue
            if not in_table:
                html_lines.append("<table>")
                # First row is header
                html_lines.append("<tr>" + "".join(f"<th>{_inline_md(c)}</th>" for c in cells) + "</tr>")
                in_table = True
            else:
                html_lines.append("<tr>" + "".join(f"<td>{_inline_md(c)}</td>" for c in cells) + "</tr>")
            continue
        else:
            if in_table:
                html_lines.append("</table>")
                in_table = False

        # Headings
        heading_match = re.match(r"^(#{1,6})\s+(.+)", line)
        if heading_match:
            level = len(heading_match.group(1))
            text = _inline_md(heading_match.group(2))
            html_lines.append(f"<h{level}>{text}</h{level}>")
            continue

        # Empty line
        if not line.strip():
            continue

        # Paragraph/text line
        html_lines.append(f"<p>{_inline_md(line)}</p>")

    if in_table:
        html_lines.append("</table>")

    html_lines.append("</body>\n</html>")
    return "\n".join(html_lines)


def _inline_md(text):
    """Convert inline markdown (bold, code, links) to HTML."""
    # Links: [text](url)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r'<a href="\2">\1</a>', text)
    # Bold: **text**
    text = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", text)
    # Inline code: `text`
    text = re.sub(r"`([^`]+)`", r"<code>\1</code>", text)
    return text


def main():
    parser = argparse.ArgumentParser(description="Generate markdown report for GenAI AIPC workflow run")
    parser.add_argument("--run-id", type=int, help="Specific workflow run ID (default: latest completed)")
    parser.add_argument("--output", "-o", type=str, help="Output file path (default: stdout)")
    parser.add_argument("--html", action="store_true", help="Also generate an HTML version of the report")
    parser.add_argument(
        "--xlsx",
        action="store_true",
        help="Generate an Excel (.xlsx) report with a single sheet and transformers version as column",
    )
    parser.add_argument(
        "--no-logs", action="store_true", help="Skip downloading logs (only show job-level conclusions)"
    )
    args = parser.parse_args()

    print("Fetching workflow run info...", file=sys.stderr)
    run_info = get_latest_run(args.run_id)
    run_id = run_info["id"]
    print(
        f"  Run #{run_info['run_number']} (ID: {run_id}), "
        f"status: {run_info.get('conclusion', run_info['status'])}",
        file=sys.stderr,
    )

    print("Fetching jobs...", file=sys.stderr)
    jobs = get_jobs(run_id)
    print(f"  Found {len(jobs)} jobs", file=sys.stderr)

    job_logs = {}
    if not args.no_logs:
        print("Downloading job logs...", file=sys.stderr)
        for i, job in enumerate(jobs):
            job_name = job["name"]
            version, device, runner = parse_job_name(job_name)
            if not version:
                continue
            print(f"  [{i+1}/{len(jobs)}] {job_name}...", end="", file=sys.stderr)
            log = download_job_log(job["id"])
            if log:
                job_logs[job["id"]] = log
                print(" OK", file=sys.stderr)
            else:
                print(" no logs", file=sys.stderr)

    print("Generating report...", file=sys.stderr)
    report = build_report(run_info, jobs, job_logs)

    if args.output:
        with open(args.output, "w", encoding="utf-8") as f:
            f.write(report)
        print(f"Report saved to {args.output}", file=sys.stderr)
    else:
        print(report)

    if args.html:
        html_content = markdown_to_html(report)
        if args.output:
            base, _ = os.path.splitext(args.output)
            html_path = base + ".html"
        else:
            html_path = "report.html"
        with open(html_path, "w", encoding="utf-8") as f:
            f.write(html_content)
        print(f"HTML report saved to {html_path}", file=sys.stderr)

    if args.xlsx:
        if args.output:
            base, _ = os.path.splitext(args.output)
            xlsx_path = base + ".xlsx"
        else:
            xlsx_path = "report.xlsx"
        build_xlsx_report(run_info, jobs, job_logs, xlsx_path)
        print(f"Excel report saved to {xlsx_path}", file=sys.stderr)


if __name__ == "__main__":
    main()
